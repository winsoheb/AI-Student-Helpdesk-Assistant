import os
import uuid
from docx import Document as DocxDocument
from database.models import User, Student, Department, Course, Subject, FAQ, Assignment, Examination, Timetable, Announcement, StudyMaterial
from werkzeug.security import generate_password_hash
from datetime import datetime, date, timedelta
from config import UPLOAD_DIR
from ai_engine.document_processor import process_and_store_document

def seed_database(db):
    if User.query.filter_by(role='admin').first():
        return
        
    print("Seeding database with realistic Indian college data and auto-generating study files...")
    
    # 1. ADMIN USER
    admin = User(username='admin', password_hash=generate_password_hash('admin123'), role='admin')
    db.session.add(admin)
    db.session.commit()
    
    # 2. DEPARTMENTS
    dept_cs = Department(department_name='Department of Computer Engineering', department_code='COMP', description='Focuses on software, hardware, and AI.')
    dept_it = Department(department_name='Department of Information Technology', department_code='IT', description='Focuses on networks, databases, and systems.')
    dept_management = Department(department_name='Department of Management Studies', department_code='MBA', description='Focuses on business administration and finance.')
    dept_entc = Department(department_name='Department of Electronics & Telecommunication', department_code='ENTC', description='Focuses on core electronics and communication networks.')
    db.session.add_all([dept_cs, dept_it, dept_management, dept_entc])
    db.session.commit()
    
    # 3. COURSES
    course_be_cs = Course(department_id=dept_cs.id, course_name='Bachelor of Engineering (Computer)', course_code='BE-COMP', duration_years=4)
    course_be_it = Course(department_id=dept_it.id, course_name='Bachelor of Engineering (IT)', course_code='BE-IT', duration_years=4)
    course_bca = Course(department_id=dept_cs.id, course_name='Bachelor of Computer Applications', course_code='BCA', duration_years=3)
    course_mba = Course(department_id=dept_management.id, course_name='Master of Business Administration', course_code='MBA-01', duration_years=2)
    course_be_entc = Course(department_id=dept_entc.id, course_name='Bachelor of Engineering (E&TC)', course_code='BE-ENTC', duration_years=4)
    db.session.add_all([course_be_cs, course_be_it, course_bca, course_mba, course_be_entc])
    db.session.commit()
    
    # 4. SUBJECTS
    sub_os = Subject(course_id=course_be_cs.id, subject_name='Operating Systems', subject_code='COMP301', semester=5)
    sub_cn = Subject(course_id=course_be_cs.id, subject_name='Computer Networks', subject_code='COMP302', semester=5)
    sub_dbms = Subject(course_id=course_be_cs.id, subject_name='Database Management Systems', subject_code='COMP303', semester=5)
    sub_ml = Subject(course_id=course_be_cs.id, subject_name='Machine Learning', subject_code='COMP401', semester=7)
    
    sub_web = Subject(course_id=course_bca.id, subject_name='Web Technologies', subject_code='BCA201', semester=3)
    sub_java = Subject(course_id=course_bca.id, subject_name='Core Java Programming', subject_code='BCA202', semester=3)
    
    sub_finance = Subject(course_id=course_mba.id, subject_name='Financial Accounting', subject_code='MBA101', semester=1)
    sub_hr = Subject(course_id=course_mba.id, subject_name='Human Resource Management', subject_code='MBA102', semester=1)
    
    db.session.add_all([sub_os, sub_cn, sub_dbms, sub_ml, sub_web, sub_java, sub_finance, sub_hr])
    db.session.commit()
    
    # 5. STUDENTS (10+ students for better scaling)
    students_data = [
        ('2021COMP001', 'Rahul Sharma', dept_cs.id, course_be_cs.id, 5, 2021),
        ('2021COMP002', 'Aditi Verma', dept_cs.id, course_be_cs.id, 5, 2021),
        ('2021COMP003', 'Karan Desai', dept_cs.id, course_be_cs.id, 5, 2021),
        ('2020COMP045', 'Neha Gupta', dept_cs.id, course_be_cs.id, 7, 2020),
        ('2022BCA012', 'Priya Patel', dept_cs.id, course_bca.id, 3, 2022),
        ('2022BCA013', 'Suresh Kumar', dept_cs.id, course_bca.id, 3, 2022),
        ('2023MBA001', 'Rohan Singh', dept_management.id, course_mba.id, 1, 2023),
        ('2023MBA002', 'Anjali Menon', dept_management.id, course_mba.id, 1, 2023),
        ('2021IT023', 'Vikram Joshi', dept_it.id, course_be_it.id, 5, 2021),
        ('2021ENTC010', 'Pooja Iyer', dept_entc.id, course_be_entc.id, 5, 2021),
    ]
    
    for roll, name, d_id, c_id, sem, yr in students_data:
        usr = User(username=roll, password_hash=generate_password_hash(roll), role='student')
        db.session.add(usr)
        db.session.commit()
        stu = Student(user_id=usr.id, student_roll_number=roll, full_name=name, department_id=d_id, course_id=c_id, semester=sem, admission_year=yr)
        db.session.add(stu)
    db.session.commit()
    
    # 6. FAQS (Expanded)
    faqs_data = [
        ('When is the library open?', 'The central library is open Monday to Saturday, 8:00 AM to 8:00 PM. Reading room is open 24/7 during exam periods.'),
        ('How do I apply for railway concession?', 'Railway concession forms are available at the Student Section (Counter No. 3) between 10 AM and 1 PM. Submit with a copy of your fee receipt and college ID.'),
        ('What is the minimum attendance requirement?', 'As per university guidelines, a minimum of 75% attendance is strictly required in both theory and practical sessions to appear for the end-semester examinations.'),
        ('Where is the training and placement cell?', 'The T&P Cell is located on the ground floor of the MBA building, Room No. 004.'),
        ('How can I pay my semester fees?', 'Fees can only be paid online through the ERP portal using Net Banking, UPI, or Debit/Credit Cards. Cash is not accepted on campus.'),
        ('Who is the HOD of Computer Engineering?', 'Dr. S. N. Patil is the current Head of the Computer Engineering Department.')
    ]
    for q, a in faqs_data:
        db.session.add(FAQ(question=q, answer=a))
    
    # 7. ASSIGNMENTS
    a1 = Assignment(title='Process Synchronization using Semaphores', description='Write a C program to solve the Producer-Consumer problem using semaphores in Linux.', subject_id=sub_os.id, issue_date=date.today(), due_date=date.today() + timedelta(days=5), maximum_marks=20)
    a2 = Assignment(title='JDBC Connectivity Project', description='Create a simple Java Swing application connected to MySQL database using JDBC.', subject_id=sub_java.id, issue_date=date.today(), due_date=date.today() + timedelta(days=12), maximum_marks=25)
    a3 = Assignment(title='Database Normalization Case Study', description='Normalize the given hospital management system database up to 3NF.', subject_id=sub_dbms.id, issue_date=date.today() - timedelta(days=2), due_date=date.today() + timedelta(days=3), maximum_marks=15)
    db.session.add_all([a1, a2, a3])
    
    # 8. EXAMINATIONS
    e1 = Examination(subject_id=sub_os.id, exam_date=date.today() + timedelta(days=20), start_time=datetime.strptime('10:00', '%H:%M').time(), end_time=datetime.strptime('12:30', '%H:%M').time(), exam_type='In-Semester Assessment', location='Block A, Room 302')
    e2 = Examination(subject_id=sub_java.id, exam_date=date.today() + timedelta(days=22), start_time=datetime.strptime('14:00', '%H:%M').time(), end_time=datetime.strptime('16:00', '%H:%M').time(), exam_type='Unit Test 2', location='Lab 4')
    e3 = Examination(subject_id=sub_dbms.id, exam_date=date.today() + timedelta(days=25), start_time=datetime.strptime('09:00', '%H:%M').time(), end_time=datetime.strptime('12:00', '%H:%M').time(), exam_type='End-Semester', location='Main Hall')
    db.session.add_all([e1, e2, e3])
    
    # 9. TIMETABLE
    t1 = Timetable(course_id=course_be_cs.id, subject_id=sub_os.id, semester=5, weekday='Monday', start_time=datetime.strptime('09:15', '%H:%M').time(), end_time=datetime.strptime('10:15', '%H:%M').time(), faculty_name='Prof. S. R. Kadam', classroom='Room 401')
    t2 = Timetable(course_id=course_be_cs.id, subject_id=sub_dbms.id, semester=5, weekday='Monday', start_time=datetime.strptime('10:15', '%H:%M').time(), end_time=datetime.strptime('11:15', '%H:%M').time(), faculty_name='Dr. A. B. Deshmukh', classroom='Room 401')
    t3 = Timetable(course_id=course_be_cs.id, subject_id=sub_cn.id, semester=5, weekday='Tuesday', start_time=datetime.strptime('11:30', '%H:%M').time(), end_time=datetime.strptime('12:30', '%H:%M').time(), faculty_name='Prof. M. K. Sharma', classroom='Room 402')
    db.session.add_all([t1, t2, t3])
    
    # 10. ANNOUNCEMENTS
    ann1 = Announcement(title='NSS Blood Donation Camp', description='The NSS unit is organizing a Blood Donation Camp on campus next Tuesday. All healthy students are encouraged to participate.', target_department_id=None, target_course_id=None)
    ann2 = Announcement(title='TCS Campus Placement Drive', description='TCS Ninja campus drive registration is open for final year BE Computer and IT students. Minimum criteria is 60% aggregate.', target_department_id=dept_cs.id, target_course_id=None)
    ann3 = Announcement(title='Fee Defaulters Notice', description='Students who have not paid the odd semester fees must pay by Friday to avoid a late fine of Rs. 500.', target_department_id=None, target_course_id=None)
    db.session.add_all([ann1, ann2, ann3])
    db.session.commit()
    
    # 11. GENERATE REAL DOCX FILES FOR AI CHATBOT TO READ
    print("Generating dummy DOCX study materials for AI testing...")
    
    # Doc 1: OS Notes
    doc1 = DocxDocument()
    doc1.add_heading('Operating Systems Notes - Module 1', 0)
    doc1.add_paragraph('An operating system (OS) is system software that manages computer hardware, software resources, and provides common services for computer programs.')
    doc1.add_paragraph('Time-sharing operating systems schedule tasks for efficient use of the system and may also include accounting software for cost allocation of processor time, mass storage, printing, and other resources.')
    doc1.add_paragraph('For hardware functions such as input and output and memory allocation, the operating system acts as an intermediary between programs and the computer hardware.')
    
    filename1 = f"{uuid.uuid4().hex}_os_notes.docx"
    filepath1 = os.path.join(UPLOAD_DIR, filename1)
    doc1.save(filepath1)
    
    mat1 = StudyMaterial(title='OS Introduction Notes', description='Basics of Operating Systems', original_filename='os_notes.docx', stored_filename=filename1, file_type='docx', file_path=filepath1, subject_id=sub_os.id, uploaded_by=admin.id)
    db.session.add(mat1)
    db.session.commit()
    process_and_store_document(mat1.id, filepath1, 'docx')
    
    # Doc 2: DBMS Notes
    doc2 = DocxDocument()
    doc2.add_heading('Database Management Systems - Normalization', 0)
    doc2.add_paragraph('Database normalization is the process of structuring a relational database in accordance with a series of so-called normal forms in order to reduce data redundancy and improve data integrity.')
    doc2.add_paragraph('Edgar F. Codd, the inventor of the relational model, introduced the concept of normalization and what we now know as the First Normal Form (1NF) in 1970.')
    doc2.add_paragraph('A relation is in Boyce-Codd Normal Form (BCNF) if and only if every determinant is a candidate key. This helps eliminate insertion, update, and deletion anomalies.')
    
    filename2 = f"{uuid.uuid4().hex}_dbms_notes.docx"
    filepath2 = os.path.join(UPLOAD_DIR, filename2)
    doc2.save(filepath2)
    
    mat2 = StudyMaterial(title='DBMS Normalization Guide', description='Normalization and BCNF', original_filename='dbms_notes.docx', stored_filename=filename2, file_type='docx', file_path=filepath2, subject_id=sub_dbms.id, uploaded_by=admin.id)
    db.session.add(mat2)
    db.session.commit()
    process_and_store_document(mat2.id, filepath2, 'docx')

    # Doc 3: Artificial Intelligence & Machine Learning Notes
    doc3 = DocxDocument()
    doc3.add_heading('Artificial Intelligence and Machine Learning - Fundamentals', 0)
    doc3.add_paragraph('Artificial Intelligence (AI) refers to the simulation of human intelligence in machines that are programmed to think like humans and mimic their actions.')
    doc3.add_paragraph('Machine Learning (ML) is a subset of AI that provides systems the ability to automatically learn and improve from experience without being explicitly programmed.')
    doc3.add_paragraph('Deep Learning is a specialized subset of Machine Learning that uses artificial neural networks with multiple layers (hence "deep") to extract higher level features from raw input.')
    doc3.add_paragraph('Supervised learning algorithms are trained using labeled examples, such as an input where the desired output is known. Unsupervised learning is used against data that has no historical labels.')
    doc3.add_paragraph('A Large Language Model (LLM) is an advanced AI algorithm that uses deep learning techniques and massive datasets to understand, summarize, generate, and predict new text content.')
    doc3.add_paragraph('Retrieval-Augmented Generation (RAG) is an AI framework that improves the quality of LLM responses by fetching facts from an external database before generating an answer. This Student Helpdesk Chatbot is a perfect example of a RAG architecture, as it searches your college PDFs to answer questions!')
    
    filename3 = f"{uuid.uuid4().hex}_ai_ml_notes.docx"
    filepath3 = os.path.join(UPLOAD_DIR, filename3)
    doc3.save(filepath3)
    
    mat3 = StudyMaterial(title='AI & Machine Learning Fundamentals', description='Intro to AI, ML, and Deep Learning', original_filename='ai_ml_notes.docx', stored_filename=filename3, file_type='docx', file_path=filepath3, subject_id=sub_ml.id, uploaded_by=admin.id)
    db.session.add(mat3)
    db.session.commit()
    process_and_store_document(mat3.id, filepath3, 'docx')

    print("Database and physical files seeded successfully.")
