import os
import subprocess
import sys

def install(package):
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

try:
    import docx
except ImportError:
    install("python-docx")
    import docx

try:
    from fpdf import FPDF
except ImportError:
    install("fpdf2")
    from fpdf import FPDF

try:
    from pptx import Presentation
except ImportError:
    install("python-pptx")
    from pptx import Presentation

docs_dir = "documentation"

def generate_docx():
    doc = docx.Document()
    doc.add_heading('AI Student Helpdesk Assistant', 0)
    doc.add_heading('Academic Project Report', 1)
    doc.add_paragraph('Submitted by: [Student Name 1], [Student Name 2]')
    doc.add_paragraph('Guide: [Guide Name]')
    
    with open(os.path.join(docs_dir, '01_PROJECT_REPORT.md'), 'r', encoding='utf-8') as f:
        lines = f.readlines()
        for line in lines:
            line = line.strip()
            if not line: continue
            if line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            else:
                doc.add_paragraph(line)
    
    doc.save(os.path.join(docs_dir, 'Complete_Project_Report.docx'))

def generate_pdf(source_md_files, output_filename, title):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", "B", 16)
    pdf.cell(w=0, h=10, text=title, new_x="LMARGIN", new_y="NEXT", align="C")
    pdf.ln(5)
    
    pdf.set_font("helvetica", size=11)
    for md_file in source_md_files:
        with open(os.path.join(docs_dir, md_file), 'r', encoding='utf-8') as f:
            for line in f:
                line = line.strip()
                if not line:
                    pdf.ln(5)
                    continue
                line = line.replace('**', '').replace('*', '')
                
                if line.startswith('## '):
                    pdf.set_font("helvetica", "B", 14)
                    pdf.cell(0, 10, line[3:], new_x="LMARGIN", new_y="NEXT")
                    pdf.set_font("helvetica", size=11)
                elif line.startswith('# '):
                    pdf.set_font("helvetica", "B", 16)
                    pdf.cell(w=0, h=10, text=line[2:], new_x="LMARGIN", new_y="NEXT")
                    pdf.set_font("helvetica", size=11)
                else:
                    try:
                        pdf.multi_cell(w=0, h=6, text=line, new_x="LMARGIN", new_y="NEXT")
                    except Exception as e:
                        print(f"Skipping line due to PDF error: {line}")
    
    pdf.output(os.path.join(docs_dir, output_filename))

def generate_pptx():
    prs = Presentation()
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "AI Student Helpdesk Assistant"
    subtitle.text = "Academic Project Presentation"
    
    bullet_slide_layout = prs.slide_layouts[1]
    
    slides_data = [
        ("Problem Statement", "Finding specific academic info in PDFs is tedious.\\nManual query processing is slow."),
        ("Objectives", "Build an offline NLP search engine.\\nCentralize academic information."),
        ("System Architecture", "Flask Backend API.\\nNext.js React Frontend.\\nSQLite Database.\\nTF-IDF Extractive QA AI."),
        ("AI Chatbot Workflow", "PDF Document Upload.\\nChunking by sentence/paragraph.\\nTF-IDF Vectorization.\\nCosine Similarity extraction.")
    ]
    
    for title_text, content_text in slides_data:
        slide = prs.slides.add_slide(bullet_slide_layout)
        shapes = slide.shapes
        title_shape = shapes.title
        body_shape = shapes.placeholders[1]
        
        title_shape.text = title_text
        tf = body_shape.text_frame
        tf.text = content_text.replace('\\n', '\n')
        
    prs.save(os.path.join(docs_dir, 'Project_Presentation.pptx'))

try:
    generate_docx()
    generate_pdf(['01_PROJECT_REPORT.md', '02_PROJECT_OVERVIEW.md'], 'Complete_Project_Report.pdf', 'Complete Project Report')
    generate_pdf(['10_USER_MANUAL.md'], 'User_Manual.pdf', 'Student User Manual')
    generate_pdf(['07_AI_CHATBOT_DOCUMENTATION.md', '08_API_DOCUMENTATION.md'], 'Technical_Documentation.pdf', 'Technical Documentation')
    generate_pdf(['13_VIVA_QUESTIONS.md'], 'Viva_Questions.pdf', 'Viva Questions and Answers')
    generate_pptx()
    print("Binaries generated successfully!")
except Exception as e:
    print(f"Error generating binaries: {e}")
